"""Pruebas del modelo de costos integrado en el tablero con un cliente de fleet-core simulado.

Los números del escenario están calculados a mano (ver el comentario ESCENARIO): las pruebas comprueban
que build_dashboard_summary y compañía reproducen esos cálculos, que el filtro de categoría no altera
el margen, que el operador no recibe costos fijos ni sueldos, y que el IVA neto convierte solo
los rubros con tarifa. No requieren el stack ni la base de datos."""
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.dashboards import export, service
from app.dashboards.schemas import DashboardSummary

D = Decimal
PERIODO = (date(2026, 9, 1), date(2026, 9, 30))

# ESCENARIO (septiembre 2026; dos vehículos, tres conductores)
#   Viajes:   d1 con v1 600 km (ingreso 750) · d1 con v2 200 km (250) · d2 con v2 200 km (250)  → 1 000 km, ingreso 1 250
#   Gastos v1: combustible 180, peaje 30, alimentación 24 (d1)   ·   mantenimiento v1 66
#   Gastos v2: combustible 120, peaje 20, alimentación 16 (d1: 60/10/8, d2: 60/10/8)  ·   mantenimiento v2 44
#   Fijos devengados: v1 depreciación 90, seguro 36, impuestos 12 · v2 depreciación 60, seguro 24, impuestos 8
#   Personal devengado: d1 sueldo 300 + beneficios 100 = 400 · d2 200 · d3 100 (no condujo)
#   Reparto del personal por km: d1 → v1 600/800 = 300 y v2 200/800 = 100; d2 → v2 200; d3 → sin asignar 100
#   v1: Cv 300  Cf 90+36+12+300 = 438  Ct 738  utilidad 750−738 = 12   margen 1,6 %   contribución 450   LP 750/300 = 2,5
#   v2: Cv 200  Cf 60+24+8+300  = 392  Ct 592  utilidad 500−592 = −92  margen −18,4 %  contribución 300   ROP −92/300
#   Flota: Cv 500  Cf 830+100 = 930  Ct 1 430  utilidad −180  margen −14,4 %  ct/km 1,43  PC 700  LP 1 250/700  ROP −180/700
V1, V2 = "aaaaaaaa-0000-0000-0000-000000000001", "aaaaaaaa-0000-0000-0000-000000000002"
D1, D2, D3 = "dddddddd-0000-0000-0000-000000000001", "dddddddd-0000-0000-0000-000000000002", "dddddddd-0000-0000-0000-000000000003"
VEH = {V1: {"id": V1, "plate": "AAA-111", "brand": "M", "model": "X", "current_mileage_km": 1000},
       V2: {"id": V2, "plate": "BBB-222", "brand": "M", "model": "Y", "current_mileage_km": 2000}}
DRIVERS = [{"id": D1, "full_name": "Ana"}, {"id": D2, "full_name": "Beto"}, {"id": D3, "full_name": "Carla"}]


def _gasto(vehicle, driver, tipo, monto, litros=0):
    return {"vehicle_id": vehicle, "driver_id": driver, "expense_date": "2026-09-10", "type": tipo,
            "amount_total": str(monto), "liters_total": str(litros), "count": 1}


def _viaje(vehicle, driver, km, ingreso):
    return {"vehicle_id": vehicle, "driver_id": driver, "trip_date": "2026-09-10", "trip_count": 1, "distance_km": km,
            "end_mileage_km": 5000, "revenue_total": str(ingreso), "revenue_trip_count": 1, "revenue_distance_km": km}


class FakeFleet:
    """Cliente de fleet-core simulado: los endpoints de costos fijos fallan si se llaman sin permiso."""

    def __init__(self, allow_fixed=True):
        self.allow_fixed, self.fixed_calls = allow_fixed, 0

    def list_vehicles(self): return list(VEH.values())
    def list_drivers(self): return list(DRIVERS)
    def list_maintenance_alerts(self): return []
    def maintenance_aggregate(self, *_):
        return [{"vehicle_id": V1, "performed_at": "2026-09-12", "category": "preventivo", "maintenance_type": "aceite", "cost_total": "66", "count": 1},
                {"vehicle_id": V2, "performed_at": "2026-09-12", "category": "preventivo", "maintenance_type": "aceite", "cost_total": "44", "count": 1}]
    def expense_aggregate(self, *_):
        return [_gasto(V1, D1, "combustible", 180, 90), _gasto(V1, D1, "peaje", 30), _gasto(V1, D1, "alimentacion", 24),
                _gasto(V2, D1, "combustible", 60, 30), _gasto(V2, D1, "peaje", 10), _gasto(V2, D1, "alimentacion", 8),
                _gasto(V2, D2, "combustible", 60, 30), _gasto(V2, D2, "peaje", 10), _gasto(V2, D2, "alimentacion", 8)]
    def trip_aggregate(self, *_):
        return [_viaje(V1, D1, 600, 750), _viaje(V2, D1, 200, 250), _viaje(V2, D2, 200, 250)]
    def fixed_cost_accrual(self, *_):
        self.fixed_calls += 1
        assert self.allow_fixed, "un rol sin permiso no debe pedir costos fijos a fleet-core"
        return [{"vehicle_id": V1, "category": "depreciacion", "amount": "90"}, {"vehicle_id": V1, "category": "seguro", "amount": "36"},
                {"vehicle_id": V1, "category": "impuestos", "amount": "12"}, {"vehicle_id": V2, "category": "depreciacion", "amount": "60"},
                {"vehicle_id": V2, "category": "seguro", "amount": "24"}, {"vehicle_id": V2, "category": "impuestos", "amount": "8"}]
    def personnel_cost_accrual(self, *_):
        self.fixed_calls += 1
        assert self.allow_fixed, "un rol sin permiso no debe pedir sueldos a fleet-core"
        return [{"driver_id": D1, "concept": "sueldo", "amount": "300"}, {"driver_id": D1, "concept": "beneficios_sociales", "amount": "100"},
                {"driver_id": D2, "concept": "sueldo", "amount": "200"}, {"driver_id": D3, "concept": "sueldo", "amount": "100"}]
    def route_aggregate(self, *_):
        return [{"origin": "Guayaquil", "destination": "Quito", "trip_count": 2, "avg_distance_km": 400, "cost_total": "400",
                 "cost_combustible": "300", "cost_peaje": "60", "cost_alimentacion": "40"}]


def _resumen(**kwargs) -> DashboardSummary:
    return service.build_dashboard_summary(FakeFleet(kwargs.pop("allow_fixed", True)), date_from=PERIODO[0], date_to=PERIODO[1], **kwargs)


def _por_placa(resumen):
    return {v.plate: v for v in resumen.vehicles}


# ------------------------------------------------------------------ vista del administrador
def test_costos_de_cada_vehiculo_con_fijos_y_personal_repartido_por_km():
    r = _resumen(include_fixed=True)
    v1, v2 = _por_placa(r)["AAA-111"].costs, _por_placa(r)["BBB-222"].costs
    assert (v1.variable_cost, v1.fixed_cost, v1.total_cost) == (D("300.00"), D("438.00"), D("738.00"))
    assert (v2.variable_cost, v2.fixed_cost, v2.total_cost) == (D("200.00"), D("392.00"), D("592.00"))
    assert (v1.personnel_cost, v2.personnel_cost) == (D("300.00"), D("300.00"))  # 300 = 400·600/800 ; 300 = 400·200/800 + 200


def test_indicadores_por_kilometro_utilidad_y_margen_del_vehiculo():
    v1 = _por_placa(_resumen(include_fixed=True))["AAA-111"].costs
    assert (v1.cv_per_km, v1.cf_per_km, v1.ct_per_km) == (D("0.5000"), D("0.7300"), D("1.2300"))
    assert (v1.revenue_per_km, v1.profit_per_km) == (D("1.2500"), D("0.0200"))
    assert (v1.profit, v1.margin_pct, v1.contribution_margin) == (D("12.00"), D("1.60"), D("450.00"))
    assert (v1.labour_productivity, v1.return_on_personnel_cost_pct) == (D("2.50"), D("4.00"))


def test_un_vehiculo_con_perdida_muestra_utilidad_y_margen_negativos():
    v2 = _por_placa(_resumen(include_fixed=True))["BBB-222"].costs
    assert (v2.profit, v2.margin_pct, v2.contribution_margin) == (D("-92.00"), D("-18.40"), D("300.00"))
    assert (v2.labour_productivity, v2.return_on_personnel_cost_pct) == (D("1.67"), D("-30.67"))


def test_costos_e_indicadores_de_la_flota_incluyen_el_personal_sin_asignar():
    r = _resumen(include_fixed=True)
    f = r.fleet_costs
    assert r.unallocated_personnel_cost == D("100.00")  # d3 cobra 100 pero no condujo
    assert (f.variable_cost, f.fixed_cost, f.total_cost) == (D("500.00"), D("930.00"), D("1430.00"))  # 830 de vehículos + 100 sin asignar
    assert (f.ct_per_km, f.profit, f.margin_pct) == (D("1.4300"), D("-180.00"), D("-14.40"))
    assert (f.personnel_cost, f.labour_productivity, f.return_on_personnel_cost_pct) == (D("700.00"), D("1.79"), D("-25.71"))
    assert f.total_cost == f.variable_cost + f.fixed_cost
    assert r.fixed_costs_omitted_reason is None and f.fixed_included


def test_la_estructura_en_ocho_categorias_suma_el_costo_total():
    f = _resumen(include_fixed=True).fleet_costs
    monto = {c.key: c.amount for c in f.composition}
    assert monto == {"combustible": D("300.00"), "reparaciones": D("110.00"), "depreciacion": D("150.00"), "capital": D("0.00"),
                     "personal": D("740.00"), "seguros_impuestos": D("80.00"), "peajes": D("50.00"), "otros": D("0.00")}
    assert sum(monto.values()) == f.total_cost
    combustible = next(c for c in f.composition if c.key == "combustible")
    assert combustible.share_pct == D("20.98") and combustible.per_km == D("0.3000")  # 300/1430 y 300/1000


def test_los_campos_heredados_del_tablero_siguen_siendo_el_gasto_operativo():
    r = _resumen(include_fixed=True)
    assert r.fleet_total_cost == D("500.00") and r.fleet_total_revenue == D("1250")
    v1 = _por_placa(r)["AAA-111"]
    assert (v1.total_cost, v1.margin, v1.revenue, v1.distance_km) == (D("300.00"), D("450.00"), D("750"), 600)
    assert (v1.breakdown.combustible, v1.breakdown.peaje, v1.breakdown.alimentacion, v1.breakdown.mantenimiento) == (D(180), D(30), D(24), D(66))
    assert r.fleet_total_margin == D("750.00")  # margen de contribución de la flota: 1 250 − 500


# ------------------------------------------------------------------ el filtro de categoría no cambia el margen
def test_el_filtro_de_categoria_acota_el_total_pero_no_el_margen_ni_los_indicadores():
    sin, con = _resumen(include_fixed=True), _resumen(include_fixed=True, category="combustible")
    assert sin.fleet_total_cost == D("500.00") and con.fleet_total_cost == D("300.00")  # el total sí se filtra
    assert con.fleet_total_margin == sin.fleet_total_margin == D("750.00")  # antes: 1 250 − 300 = 950
    assert con.fleet_costs == sin.fleet_costs
    assert _por_placa(con)["AAA-111"].margin == _por_placa(sin)["AAA-111"].margin == D("450.00")
    assert _por_placa(con)["AAA-111"].costs == _por_placa(sin)["AAA-111"].costs


# ------------------------------------------------------------------ rol operador y período
def test_el_operador_no_recibe_costos_fijos_ni_sueldos_y_no_se_le_finge_un_total():
    fleet = FakeFleet(allow_fixed=False)
    r = service.build_dashboard_summary(fleet, date_from=PERIODO[0], date_to=PERIODO[1], include_fixed=False)
    assert fleet.fixed_calls == 0
    f = r.fleet_costs
    assert (f.fixed_included, r.fixed_costs_omitted_reason, r.unallocated_personnel_cost) == (False, "rol", None)
    assert (f.variable_cost, f.cv_per_km, f.contribution_margin) == (D("500.00"), D("0.5000"), D("750.00"))
    for campo in ("fixed_cost", "total_cost", "ct_per_km", "profit", "margin_pct", "personnel_cost", "labour_productivity"):
        assert getattr(f, campo) is None, campo
    assert all(c.share_pct is not None for c in f.composition if c.key in ("combustible", "peajes"))
    assert sum(c.share_pct for c in f.composition) == D("100.00")  # sobre lo que ve: costos variables


def test_sin_periodo_no_se_prorratean_costos_fijos():
    fleet = FakeFleet()
    r = service.build_dashboard_summary(fleet, include_fixed=True)  # sin fechas
    assert fleet.fixed_calls == 0 and r.fixed_costs_omitted_reason == "sin_periodo" and not r.fleet_costs.fixed_included


def test_un_vehiculo_elegido_usa_los_km_de_toda_la_flota_para_repartir_el_personal():
    r = _resumen(include_fixed=True, vehicle_id=V1)
    assert [v.plate for v in r.vehicles] == ["AAA-111"]
    assert r.vehicles[0].costs.personnel_cost == D("300.00")  # igual que en la vista de la flota
    assert r.unallocated_personnel_cost is None  # lo que no es de ningún vehículo no se muestra
    assert r.fleet_costs.total_cost == r.vehicles[0].costs.total_cost == D("738.00")


# ------------------------------------------------------------------ IVA neto
def test_iva_neto_convierte_solo_los_rubros_con_tarifa_y_no_toca_los_ingresos():
    bruto, neto = _resumen(include_fixed=True, iva_mode="bruto"), _resumen(include_fixed=True, iva_mode="neto")
    assert (bruto.iva_mode, neto.iva_mode) == ("bruto", "neto")
    v1 = _por_placa(neto)["AAA-111"]
    # combustible 180, alimentación 24 y mantenimiento 66 llevan 15 %; el peaje no
    assert (v1.breakdown.combustible, v1.breakdown.peaje) == (D("156.52"), D("30.00"))  # 180/1,15
    assert (v1.breakdown.alimentacion, v1.breakdown.mantenimiento) == (D("20.87"), D("57.39"))  # 24/1,15 ; 66/1,15
    assert v1.costs.variable_cost == D("264.78")  # 270/1,15 + 30
    assert v1.revenue == D("750") and v1.costs.revenue == D("750.00")
    assert neto.fleet_costs.fixed_cost == bruto.fleet_costs.fixed_cost == D("930.00")  # fijos y personal: tarifa 0
    assert neto.fleet_costs.variable_cost < bruto.fleet_costs.variable_cost


def test_un_modo_de_iva_desconocido_se_rechaza():
    with pytest.raises(ValueError, match="iva_mode"):
        _resumen(iva_mode="gravado")


def test_por_defecto_el_iva_es_bruto():
    assert _resumen().iva_mode == "bruto"


# ------------------------------------------------------------------ conductores
def test_el_admin_ve_el_personal_del_conductor_y_su_productividad_laboral():
    r = service.build_driver_summary(FakeFleet(), *PERIODO, include_fixed=True)
    por = {d.full_name: d for d in r.drivers}
    assert r.personnel_included is True
    assert (por["Ana"].personnel_cost, por["Ana"].labour_productivity) == (D("400.00"), D("2.50"))  # 1 000 / 400
    assert (por["Beto"].personnel_cost, por["Beto"].labour_productivity) == (D("200.00"), D("1.25"))  # 250 / 200
    assert (por["Carla"].personnel_cost, por["Carla"].labour_productivity) == (D("100.00"), None)  # sin ingresos
    assert (por["Ana"].total_cost, por["Ana"].margin, por["Beto"].margin) == (D("312.00"), D("688.00"), D("172.00"))


def test_el_operador_no_ve_sueldos_por_conductor():
    fleet = FakeFleet(allow_fixed=False)
    r = service.build_driver_summary(fleet, *PERIODO, include_fixed=False)
    assert fleet.fixed_calls == 0 and r.personnel_included is False
    assert all(d.personnel_cost is None and d.labour_productivity is None for d in r.drivers)
    assert {d.full_name: d.margin for d in r.drivers}["Ana"] == D("688.00")


def test_el_margen_del_conductor_no_depende_del_filtro_de_categoria():
    con = {d.full_name: d for d in service.build_driver_summary(FakeFleet(), *PERIODO, category="combustible").drivers}
    assert con["Ana"].total_cost == D("240.00") and con["Ana"].margin == D("688.00")  # antes: 1 000 − 240


# ------------------------------------------------------------------ rutas, tendencia y registros con IVA
def test_las_rutas_convierten_cada_tipo_de_gasto_en_modo_neto():
    bruto = service.build_route_summary(FakeFleet(), *PERIODO, iva_mode="bruto").routes[0]
    neto = service.build_route_summary(FakeFleet(), *PERIODO, iva_mode="neto").routes[0]
    assert bruto.cost_total == D("400.00")
    assert neto.cost_total == D("355.65")  # 340/1,15 + 60
    assert neto.cost_per_km < bruto.cost_per_km


def test_la_tendencia_y_los_registros_respetan_el_modo_de_iva():
    t = service.build_cost_trend(FakeFleet(), "month", None, *PERIODO, iva_mode="neto")
    assert t.iva_mode == "neto" and t.points[0].breakdown.mantenimiento == D("95.65")  # (66 + 44) / 1,15
    assert t.points[0].period_start == date(2026, 9, 1)  # inicio del período del punto (mes)
    registros = service.build_expense_records(FakeFleet(), None, *PERIODO, category="alimentacion", iva_mode="neto")
    assert {r.amount for r in registros} == {D("20.87"), D("6.96")}  # 24/1,15 y 8/1,15


# ------------------------------------------------------------------ exportación
def test_la_exportacion_del_admin_incluye_costos_fijos_estructura_y_notas():
    r = _resumen(include_fixed=True, iva_mode="neto")
    wb = load_workbook(BytesIO(export.build_costs_workbook(r)))
    assert wb.sheetnames == ["Costos por vehiculo", "Estructura de costos", "Notas"]
    cabecera = [c.value for c in wb["Costos por vehiculo"][1]]
    assert "Costo total (Ct)" in cabecera and "Margen %" in cabecera
    filas = list(wb["Costos por vehiculo"].iter_rows(values_only=True))
    total = next(f for f in filas if f[0] == "TOTAL FLOTA")
    assert total[cabecera.index("Costo total (Ct)")] == pytest.approx(float(r.fleet_costs.total_cost))
    assert len(list(wb["Estructura de costos"].iter_rows())) == 9  # cabecera + 8 categorías
    notas = " ".join(str(c) for fila in wb["Notas"].iter_rows(values_only=True) for c in fila if c)
    assert "sin IVA (neto)" in notas and "Poliak" in notas and "Zysińska" in notas and "Kutac" in notas
    assert export.build_costs_pdf(r)[:4] == b"%PDF"


def test_la_exportacion_sin_costos_fijos_no_finge_un_costo_total():
    r = service.build_dashboard_summary(FakeFleet(allow_fixed=False), date_from=PERIODO[0], date_to=PERIODO[1])
    wb = load_workbook(BytesIO(export.build_costs_workbook(r)))
    cabecera = [c.value for c in wb["Costos por vehiculo"][1]]
    assert "Costo total (Ct)" not in cabecera and "Margen de contribución" in cabecera
    assert "no incluidos" in " ".join(str(c) for fila in wb["Notas"].iter_rows(values_only=True) for c in fila if c)
    assert export.build_costs_pdf(r)[:4] == b"%PDF"
