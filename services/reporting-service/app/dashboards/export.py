"""Exportación del resumen de costos a Excel y PDF (RF-17).

Solo la piden administradores, así que pueden incluir los costos fijos y de personal (Fase 3): cuando el
resumen los trae, se agregan Cf, Ct, ingreso, utilidad y margen, y la estructura del costo en las ocho
categorías de Zysińska & Menes (2026). El modo de IVA y la fórmula de cada indicador viajan en el archivo.
"""
from decimal import Decimal
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.dashboards import cost_model
from app.dashboards.schemas import DashboardSummary, VehicleSummary


def _f(value: Optional[Decimal]) -> Optional[float]:
    return None if value is None else float(value)


def _fixed_included(summary: DashboardSummary) -> bool:
    return bool(summary.fleet_costs and summary.fleet_costs.fixed_included)


def _iva_label(summary: DashboardSummary) -> str:
    return "sin IVA (neto)" if summary.iva_mode == "neto" else "con IVA (bruto: montos pagados)"


def _vehicle_row(v: VehicleSummary, fixed: bool) -> list:
    c = v.costs
    base = [
        v.plate, v.brand, v.model, v.current_mileage_km, v.trip_count, v.distance_km,
        float(v.breakdown.combustible), float(v.breakdown.peaje), float(v.breakdown.alimentacion),
        float(v.breakdown.mantenimiento), _f(c.variable_cost) if c else float(v.total_cost),
    ]
    if fixed and c:
        base += [_f(c.fixed_cost), _f(c.personnel_cost), _f(c.total_cost), _f(c.ct_per_km), _f(c.revenue), _f(c.profit), _f(c.margin_pct)]
    else:
        base += [_f(c.cv_per_km) if c else None, _f(c.revenue) if c else None, _f(c.contribution_margin) if c else None]
    return base + [v.pending_maintenance_alerts]


def _headers(fixed: bool) -> list[str]:
    base = ["Placa", "Marca", "Modelo", "Km actual", "Viajes", "Km recorridos", "Combustible", "Peajes",
            "Alimentación", "Mantenimiento", "Costo variable (Cv)"]
    if fixed:
        base += ["Costo fijo (Cf)", "del cual personal", "Costo total (Ct)", "Costo/km total (ct)", "Ingreso", "Utilidad", "Margen %"]
    else:
        base += ["Costo/km variable (cv)", "Ingreso", "Margen de contribución"]
    return base + ["Alertas de mantenimiento"]


def build_costs_workbook(summary: DashboardSummary) -> bytes:
    """RF-17: reporte de costos y mantenimientos por vehículo, exportable a Excel."""
    fixed = _fixed_included(summary)
    wb = Workbook()
    ws = wb.active
    ws.title = "Costos por vehiculo"

    ws.append(_headers(fixed))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for v in summary.vehicles:
        ws.append(_vehicle_row(v, fixed))

    ws.append([])
    fleet = summary.fleet_costs
    total_row = ["TOTAL FLOTA", "", "", "", "", fleet.distance_km if fleet else "", "", "", "", "", _f(fleet.variable_cost) if fleet else float(summary.fleet_total_cost)]
    if fixed and fleet:
        total_row += [_f(fleet.fixed_cost), _f(fleet.personnel_cost), _f(fleet.total_cost), _f(fleet.ct_per_km), _f(fleet.revenue), _f(fleet.profit), _f(fleet.margin_pct)]
    else:
        total_row += [_f(fleet.cv_per_km) if fleet else None, _f(fleet.revenue) if fleet else None, _f(fleet.contribution_margin) if fleet else None]
    ws.append(total_row + [""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 28)

    # Estructura del costo en las ocho categorías (Zysińska & Menes, 2026): monto por vehículo y en la flota.
    if summary.vehicles and summary.vehicles[0].costs:
        st = wb.create_sheet("Estructura de costos")
        st.append(["Categoría (Zysińska & Menes, 2026)"] + [v.plate for v in summary.vehicles] + ["Flota", "% de la flota", "Costo/km flota"])
        for cell in st[1]:
            cell.font = Font(bold=True)
        for index, (key, label) in enumerate(cost_model.CATEGORIES):
            line = fleet.composition[index] if fleet else None
            st.append(
                [label]
                + [_f(v.costs.composition[index].amount) for v in summary.vehicles]
                + [_f(line.amount) if line else None, _f(line.share_pct) if line else None, _f(line.per_km) if line else None]
            )
        st.column_dimensions["A"].width = 58

    # Notas: qué incluye el archivo y de dónde sale cada indicador.
    notes = wb.create_sheet("Notas")
    notes.append(["Modo de IVA", _iva_label(summary)])
    notes.append(["Costos fijos y de personal", "incluidos" if fixed else "no incluidos (falta el período o el rol no los ve)"])
    if summary.unallocated_personnel_cost:
        notes.append(["Personal sin kilómetros en el período (solo en el total de la flota)", float(summary.unallocated_personnel_cost)])
    if fleet and fleet.revenue_coverage_pct is not None and fleet.revenue_coverage_pct < 100:
        notes.append(["Cobertura de ingresos (% de km con ingreso registrado)", float(fleet.revenue_coverage_pct)])
    notes.append([])
    notes.append(["Indicador", "Fórmula", "Fuente"])
    for cell in notes[notes.max_row]:
        cell.font = Font(bold=True)
    for ind in cost_model.INDICATORS:
        notes.append([f"{ind.id} — {ind.name}", ind.formula, ind.source])
    notes.column_dimensions["A"].width = 60
    notes.column_dimensions["B"].width = 80
    notes.column_dimensions["C"].width = 60

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_costs_pdf(summary: DashboardSummary) -> bytes:
    """RF-17: mismo reporte de costos por vehículo, exportado a PDF (horizontal, con Cf y Ct si se incluyen)."""
    fixed = _fixed_included(summary)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), title="Reporte de costos - FlotaControl")
    styles = getSampleStyleSheet()
    fleet = summary.fleet_costs

    story = [
        Paragraph("FlotaControl — Reporte de costos por vehículo", styles["Title"]),
        Paragraph(
            f"Montos {_iva_label(summary)}. Costos fijos y de personal: "
            f"{'incluidos (Ct = Cv + Cf, Poliak et al., 2025)' if fixed else 'no incluidos'}.",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]

    headers = ["Placa", "Marca/Modelo", "Km", "Viajes", "Combustible", "Peajes", "Alimentación", "Mantenim.", "Cv"]
    if fixed:
        headers += ["Cf", "Ct", "ct/km", "Ingreso", "Utilidad", "Margen %"]
    else:
        headers += ["cv/km", "Ingreso", "Contribución"]
    data = [headers + ["Alertas"]]

    def money(x: Optional[Decimal]) -> str:
        return "—" if x is None else f"{x:.2f}"

    for v in summary.vehicles:
        c = v.costs
        row = [v.plate, f"{v.brand} {v.model}", str(v.distance_km), str(v.trip_count),
               money(v.breakdown.combustible), money(v.breakdown.peaje), money(v.breakdown.alimentacion),
               money(v.breakdown.mantenimiento), money(c.variable_cost if c else v.total_cost)]
        if fixed and c:
            row += [money(c.fixed_cost), money(c.total_cost), "—" if c.ct_per_km is None else f"{c.ct_per_km:.3f}",
                    money(c.revenue), money(c.profit), money(c.margin_pct)]
        else:
            row += ["—" if not c or c.cv_per_km is None else f"{c.cv_per_km:.3f}", money(c.revenue if c else None), money(c.contribution_margin if c else None)]
        data.append(row + [str(v.pending_maintenance_alerts)])

    total = ["TOTAL FLOTA", "", str(fleet.distance_km) if fleet else "", "", "", "", "", "", money(fleet.variable_cost if fleet else summary.fleet_total_cost)]
    if fixed and fleet:
        total += [money(fleet.fixed_cost), money(fleet.total_cost), "—" if fleet.ct_per_km is None else f"{fleet.ct_per_km:.3f}",
                  money(fleet.revenue), money(fleet.profit), money(fleet.margin_pct)]
    else:
        total += ["—" if not fleet or fleet.cv_per_km is None else f"{fleet.cv_per_km:.3f}", money(fleet.revenue if fleet else None), money(fleet.contribution_margin if fleet else None)]
    data.append(total + [""])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14283f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#eef1f4")]),
            ]
        )
    )
    story.append(table)

    if fixed and fleet:
        story += [Spacer(1, 12), Paragraph("Estructura del costo por kilómetro (Zysińska & Menes, 2026)", styles["Heading3"])]
        rows = [["Categoría", "Monto", "Costo/km", "% del total"]]
        for line in fleet.composition:
            rows.append([line.label, money(line.amount), "—" if line.per_km is None else f"{line.per_km:.4f}", "—" if line.share_pct is None else f"{line.share_pct:.1f}%"])
        composition = Table(rows, repeatRows=1, colWidths=[300, 80, 80, 80])
        composition.setStyle(
            TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14283f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.5, colors.grey)])
        )
        story.append(composition)

    doc.build(story)
    return buffer.getvalue()
